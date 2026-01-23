import sys
import os
from flask import Flask, Response, render_template, request, jsonify
import time, json, serial, threading
from openpyxl import Workbook, load_workbook
from datetime import datetime

# ================= FLASK CONFIG FOR EXE =================
# This block fixes the TemplateNotFound error
if getattr(sys, 'frozen', False):
    # If running as an exe, use the temporary folder where PyInstaller unpacks
    template_folder = os.path.join(sys._MEIPASS, 'templates')
    app = Flask(__name__, template_folder=template_folder)
else:
    # If running as a script, use the default folder
    app = Flask(__name__)

# ================= SERIAL CONFIG =================
SERIAL_PORT = "COM"+str(input("Enter COM port: "))
BAUD_RATE   = 115200

# Wrap serial connection in try/except so app opens even if COM is missing
try:
    ser = serial.Serial(SERIAL_PORT, BAUD_RATE, timeout=0.1)
    print(f"Connected to {SERIAL_PORT}")
except Exception as e:
    ser = None
    print(f"Warning: Could not connect to {SERIAL_PORT}. Error: {e}")

# ================= GLOBAL STATE =================
latest_data = {"C0": 0, "C1": 0, "C2": 0, "C3": 0}
excel_file = None
lock = threading.Lock()

# ================= SERIAL READER THREAD =================
def serial_reader():
    global latest_data
    if ser is None: 
        return # Stop thread if serial failed

    buffer = ""

    while True:
        try:
            if ser.in_waiting:
                buffer += ser.read(ser.in_waiting).decode(errors="ignore")

                while "\n" in buffer:
                    line, buffer = buffer.split("\n", 1)
                    line = line.strip()

                    # Expected: C0=123 C1=456 C2=789 C3=321
                    values = {}

                    for token in line.split():
                        if "=" in token:
                            k, v = token.split("=", 1)
                            if k in latest_data:
                                try:
                                    values[k] = float(v)
                                except ValueError:
                                    pass

                    if len(values) == 4:
                        with lock:
                            latest_data = values.copy()

        except Exception as e:
            print("Serial error:", e)

        time.sleep(0.001)

# ================= START SERIAL THREAD =================
threading.Thread(target=serial_reader, daemon=True).start()

# ================= STREAM TO BROWSER =================
def stream_data():
    while True:
        with lock:
            data = latest_data.copy()

        yield f"data:{json.dumps(data)}\n\n"
        time.sleep(0.02)   # ~50 FPS browser-friendly

@app.route("/data")
def data():
    return Response(stream_data(), mimetype="text/event-stream")

@app.route("/")
def index():
    # Make sure 'indextry.html' is inside a folder named 'templates'
    return render_template("indextry.html")

# ================= CREATE EXCEL =================
@app.route("/create_excel", methods=["POST"])
def create_excel():
    global excel_file
    
    # Save file in the same folder as the EXE, not the temp folder
    if getattr(sys, 'frozen', False):
        application_path = os.path.dirname(sys.executable)
    else:
        application_path = os.path.dirname(os.path.abspath(__file__))

    filename = f"hand_measurements_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    excel_file = os.path.join(application_path, filename)

    wb = Workbook()
    ws = wb.active
    ws.append([
        "Timestamp", "Name", "DOB", "Age", "Description",
        "Movement", "Finger", "C0", "C1", "C2", "C3"
    ])
    wb.save(excel_file)

    return jsonify({"message": f"Excel created: {filename}"})

# ================= CAPTURE (NON-BLOCKING) =================
@app.route("/capture", methods=["POST"])
def capture():
    if not excel_file:
        return jsonify({"error": "Create Excel first"}), 400

    data = request.get_json(silent=True) or {}

    threading.Thread(
        target=save_measurements,
        args=(data,),
        daemon=True
    ).start()

    return jsonify({"message": "Capture started (non-blocking)"})

# ================= BACKGROUND SAVE =================
def save_measurements(data):
    wb = load_workbook(excel_file)
    ws = wb.active

    name   = data.get("name", "")
    dob    = data.get("dob", "")
    age    = data.get("age", "")
    desc   = data.get("description", "")
    move   = data.get("movement", "")
    finger = data.get("finger", "")

    for _ in range(10):
        with lock:
            d = latest_data.copy()

        if move == "Lateral":
            row = [
                datetime.now().isoformat(),
                name, dob, age, desc, move, finger,
                d["C0"], d["C1"], d["C2"], "-"
            ]
        else:
            row = [
                datetime.now().isoformat(),
                name, dob, age, desc, move, finger,
                "-", "-", "-", d["C3"]
            ]

        ws.append(row)
        time.sleep(0.5)

    wb.save(excel_file)

# ================= RUN =================
if __name__ == "__main__":
    # use_reloader=False is important for threading in exe
    app.run(debug=False, threaded=True, use_reloader=False)